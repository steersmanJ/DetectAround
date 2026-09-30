import io
import datetime
import requests
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from vworld_api import get_detailed_address, get_static_map_image
from gee_analyzer import get_yearly_satellite_images, get_ndvi_timeseries

def create_ndvi_graph(ts_data):
    """
    NDVI 시계열 데이터를 바탕으로 matplotlib 그래프 이미지를 생성하여 바이트로 반환합니다.
    """
    if not ts_data:
        return None
        
    dates = [x[0] for x in ts_data]
    ndvis = [x[1] for x in ts_data]
    
    # 폰트 깨짐 방지를 위해 영어로 설정하거나, 폰트 설정을 스킵하고 기본 폰트 사용
    plt.figure(figsize=(7, 3))
    plt.plot(dates, ndvis, marker='o', linestyle='-', color='g')
    plt.axhline(y=0.25, color='r', linestyle='--', label='Threshold (0.25)')
    
    plt.title('NDVI Time-Series (Last 3 Years)', fontsize=12)
    plt.xlabel('Date')
    plt.ylabel('NDVI (Vegetation Index)')
    plt.grid(True, linestyle=':', alpha=0.6)
    
    # x축 레이블 간격 조절 (너무 촘촘하지 않게)
    plt.xticks(dates[::max(1, len(dates)//6)], rotation=20)
    plt.legend()
    plt.tight_layout()
    
    buf = io.BytesIO()
    plt.savefig(buf, format='png', dpi=100)
    plt.close()
    buf.seek(0)
    return buf

def generate_report_docx(lon, lat, pnu, jibun_short, jimok_char, jimok_desc, ndvi_score, jiga=None):
    """
    python-docx를 사용하여 농지법 위반 의심 현장 조사 보고서를 생성합니다.
    """
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(11)
    
    # ── 제목 ──
    title = doc.add_heading('농지법 위반 의심 현장 조사 보고서', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # ── 1. 대상 농지 정보 ──
    doc.add_heading('[ 1. 대상 농지 정보 ]', level=2)
    addr_info = get_detailed_address(lon, lat)
    parcel_addr = addr_info.get("parcel") or jibun_short
    road_addr = addr_info.get("road") or "(도로명 주소 없음)"
    
    p_info = doc.add_paragraph()
    p_info.add_run("고유번호(PNU) : ").bold = True
    p_info.add_run(f"{pnu}\n")
    p_info.add_run("지번 주소 : ").bold = True
    p_info.add_run(f"{parcel_addr}\n")
    p_info.add_run("도로명 주소 : ").bold = True
    p_info.add_run(f"{road_addr}\n")
    p_info.add_run("지목 : ").bold = True
    p_info.add_run(f"{jimok_char} ({jimok_desc})\n")
    
    if jiga:
        try:
            formatted_jiga = f"{int(jiga):,}원/㎡"
            p_info.add_run("개별공시지가 : ").bold = True
            p_info.add_run(f"{formatted_jiga}\n")
        except:
            pass
            
    p_info.add_run("조사 일자 : ").bold = True
    p_info.add_run(f"{datetime.datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph()

    # ── 2. 현장 위성 및 지적도 사진 ──
    doc.add_heading('[ 2. 현재 상태 (일반지도 및 지적위성) ]', level=2)
    graphic_img = get_static_map_image(lon, lat, map_type="GRAPHIC")
    hybrid_img = get_static_map_image(lon, lat, map_type="PHOTO_HYBRID")
    
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    if graphic_img:
        run = table.cell(0, 0).paragraphs[0].add_run()
        run.add_picture(io.BytesIO(graphic_img), width=Inches(2.8))
        table.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    if hybrid_img:
        run2 = table.cell(0, 1).paragraphs[0].add_run()
        run2.add_picture(io.BytesIO(hybrid_img), width=Inches(2.8))
        table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    
    # ── 2-1. 카카오 초고해상도 스카이뷰 (RPA 캡처) ──
    try:
        from kakao_skyview_bot import capture_skyview
        sv_bytes = capture_skyview(lat, lon)
        if sv_bytes:
            p_sv_desc = doc.add_paragraph("▼ 초고해상도 스카이뷰 (지적편집도 적용, 최대 확대)")
            p_sv_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sv = doc.add_paragraph()
            p_sv.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sv.add_run().add_picture(io.BytesIO(sv_bytes), width=Inches(6.0))
    except Exception as e:
        doc.add_paragraph(f"(스카이뷰 RPA 구동 실패: {e})")
    doc.add_paragraph()

    # ── 3. 식생지수 시계열 분석 (과거 3년) ──
    doc.add_heading('[ 3. 과거 3년 식생 변화(NDVI) 추이 ]', level=2)
    doc.add_paragraph("아래 그래프는 과거 3년간 해당 토지의 식물 활력도(NDVI) 변화를 나타냅니다. 0.25 이하로 떨어진 시점부터 식물이 사라진(불법 전용된) 것으로 강력히 의심됩니다.")
    
    ts_data = get_ndvi_timeseries(lon, lat, years=3)
    graph_buf = create_ndvi_graph(ts_data)
    if graph_buf:
        p_graph = doc.add_paragraph()
        p_graph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_graph.add_run().add_picture(graph_buf, width=Inches(6.0))
    else:
        doc.add_paragraph("(시계열 데이터를 가져오지 못했습니다.)")
    doc.add_paragraph()

    # ── 4. 과거 3년 여름철 위성사진 비교 ──
    doc.add_heading('[ 4. 과거 3년 위성사진 비교 ]', level=2)
    doc.add_paragraph("연도별(6~8월) 맑은 날을 합성한 위성사진입니다. 대상 농지의 변화를 시각적으로 확인할 수 있습니다.")
    
    yearly_images = get_yearly_satellite_images(lon, lat, years=3)
    if yearly_images:
        num_imgs = len(yearly_images)
        t_img = doc.add_table(rows=1, cols=num_imgs)
        t_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, item in enumerate(yearly_images):
            cell = t_img.cell(0, i)
            url = item.get('url')
            year = item.get('year')
            p_cell = cell.paragraphs[0]
            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if url:
                try:
                    res = requests.get(url, verify=False)
                    if res.status_code == 200:
                        p_cell.add_run().add_picture(io.BytesIO(res.content), width=Inches(2.0))
                except:
                    p_cell.add_run("(이미지 로드 실패)")
            else:
                p_cell.add_run("(구름으로 인해 사진 없음)")
            cell.add_paragraph(f"{year}년 여름").alignment = WD_ALIGN_PARAGRAPH.CENTER
            
    doc.add_paragraph()
    doc.add_paragraph("▼ [RPA 초고해상도 캡처] 카카오맵 과거 위성사진 비교")
    
    try:
        from kakao_skyview_bot import capture_historical_skyviews
        hist_skyviews = capture_historical_skyviews(lat, lon, num_years=3)
        
        if hist_skyviews:
            t_hist = doc.add_table(rows=1, cols=len(hist_skyviews))
            t_hist.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for i, item in enumerate(hist_skyviews):
                cell = t_hist.cell(0, i)
                p_cell = cell.paragraphs[0]
                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cell.add_run().add_picture(io.BytesIO(item['bytes']), width=Inches(2.0))
                cell.add_paragraph(f"{item['year']} (카카오)").alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:
        doc.add_paragraph(f"(카카오 과거 위성사진 캡처 실패: {e})")
        
    doc.add_paragraph()
    # ── 5. 현장 로드뷰 (RPA 캡처) ──
    doc.add_heading('[ 5. 카카오맵 현장 로드뷰 ]', level=2)
    doc.add_paragraph("파이썬 RPA를 통해 자동 캡처된 현장 주변의 로드뷰 사진입니다. (지형지물 및 불법 건축물/주차장 확인용)")
    
    try:
        from roadview_bot import capture_roadview
        rv_bytes_list = capture_roadview(lat, lon)
        if rv_bytes_list and len(rv_bytes_list) > 0:
            if len(rv_bytes_list) == 1:
                p_rv = doc.add_paragraph()
                p_rv.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_rv.add_run().add_picture(io.BytesIO(rv_bytes_list[0]), width=Inches(6.0))
            else:
                doc.add_paragraph("💡 아래 4장의 360도 회전 사진 중 타겟 농지가 가장 잘 보이는 사진만 남기고 나머지는 삭제해 주세요.")
                rv_table = doc.add_table(rows=2, cols=2)
                rv_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                for idx, img_bytes in enumerate(rv_bytes_list):
                    row_idx = idx // 2
                    col_idx = idx % 2
                    cell = rv_table.cell(row_idx, col_idx)
                    p_cell = cell.paragraphs[0]
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cell.add_run().add_picture(io.BytesIO(img_bytes), width=Inches(2.8))
                    cell.add_paragraph(f"사진 {idx+1}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_paragraph("(로드뷰를 캡처하지 못했거나 해당 지역에 로드뷰가 없습니다.)")
    except Exception as e:
        doc.add_paragraph(f"(RPA 모듈 구동 실패: {e})")
    doc.add_paragraph()

    # ── 6. 국민신문고 신고용 텍스트 ──
    doc.add_heading('[ 6. 국민신문고 신고 템플릿 ]', level=2)
    template_text = (
        f"농지법 위반 신고합니다.\n"
        f"대상 지번: {parcel_addr} (지목: {jimok_desc})\n\n"
        f"위 토지는 농지임에도 불구하고, 위성 분석 결과 식생이 오랫동안 확인되지 않고 있습니다. "
        f"보고서에 첨부된 과거 3년간의 위성사진과 식생지수(NDVI) 그래프를 확인해 보시면, "
        f"특정 시점부터 식물이 완전히 사라진 것을 명확히 알 수 있습니다.\n\n"
        f"현재 무단 전용(건축물, 주차장, 적치장 등)되어 사용되고 있을 확률이 매우 높으므로, "
        f"신속한 현장 점검과 원상회복 명령 처분을 요청드립니다."
    )
    p_template = doc.add_paragraph(template_text)
    p_template.paragraph_format.left_indent = Inches(0.5)
    
    doc.add_paragraph()
    p_footer = doc.add_paragraph("본 보고서는 Global Vegetation Analysis Dashboard에 의해 자동 생성되었습니다.")
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    
    return doc_io.read()
